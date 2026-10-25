from openpyxl import Workbook, load_workbook
import os

ARCHIVO = "datos.xlsx"


def obtenerLibro():
    if os.path.exists(ARCHIVO):
        return load_workbook(ARCHIVO)

    return Workbook()


def guardar(hoja):
    libro = hoja.parent
    libro.save(ARCHIVO)